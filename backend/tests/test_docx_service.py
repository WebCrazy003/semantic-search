# backend/tests/test_docx_service.py
from pathlib import Path

import pytest

from app.services.docx_service import DocxService
from app.services.extractors import ExtractionError, ExtractionUnsupportedError


@pytest.fixture
def service() -> DocxService:
    return DocxService(min_document_chars=20, extract_tables=True)


def _extract(service: DocxService, path: Path):
    return service.extract(path, document_id="doc", file_hash="doc")


class TestStructure:
    def test_headings_and_paragraphs_keep_document_order(
        self, service: DocxService, docx_dir: Path
    ) -> None:
        first = _extract(service, docx_dir / "manual_ko.docx").pages[0]
        assert [block.kind for block in first.blocks] == [
            "heading",
            "paragraph",
            "paragraph",
            "paragraph",
        ]
        assert first.blocks[0].text == "제1장 안전 주의사항"
        assert [block.order for block in first.blocks] == [0, 1, 2, 3]

    def test_a_heading_style_with_a_korean_name_is_found_by_its_outline_level(
        self, service: DocxService, docx_dir: Path
    ) -> None:
        second = _extract(service, docx_dir / "manual_ko.docx").pages[1]
        assert second.blocks[0].kind == "heading"
        assert second.blocks[0].text == "제2장 정기 점검"

    def test_metadata_and_language(self, service: DocxService, docx_dir: Path) -> None:
        document = _extract(service, docx_dir / "manual_ko.docx")
        assert document.meta.title == "장비 사용 설명서"
        assert document.meta.language == "ko"
        assert document.meta.filename == "manual_ko.docx"

    def test_a_table_becomes_one_markdown_block(self, service: DocxService, docx_dir: Path) -> None:
        blocks = _extract(service, docx_dir / "maintenance_zh.docx").pages[0].blocks
        tables = [block for block in blocks if block.kind == "table"]
        assert len(tables) == 1
        rows = tables[0].text.split("\n")
        assert rows[0] == "| 部件 / 부품 | 周期 / 주기 | 型号 / 모델 |"
        assert rows[1] == "| --- | --- | --- |"

    def test_a_vertically_merged_cell_is_not_repeated(
        self, service: DocxService, docx_dir: Path
    ) -> None:
        table = next(
            block
            for block in _extract(service, docx_dir / "maintenance_zh.docx").pages[0].blocks
            if block.kind == "table"
        )
        assert table.text.count("每年 / 매년") == 1

    def test_tables_can_be_switched_off(self, docx_dir: Path) -> None:
        service = DocxService(extract_tables=False)
        blocks = _extract(service, docx_dir / "maintenance_zh.docx").pages[0].blocks
        assert "table" not in {block.kind for block in blocks}
        assert any(block.text == "XJ-200B" for block in blocks)


class TestText:
    def test_soft_and_non_breaking_hyphens(self, service: DocxService, docx_dir: Path) -> None:
        text = _extract(service, docx_dir / "maintenance_zh.docx").pages[0].text
        assert "maintenance kit, model XJ-200B." in text

    def test_line_breaks_inside_chinese_add_no_space(
        self, service: DocxService, docx_dir: Path
    ) -> None:
        text = _extract(service, docx_dir / "maintenance_zh.docx").pages[0].text
        assert "更换滤芯之前必须关闭电源。" in text

    def test_tracked_insertions_are_kept_and_deletions_dropped(
        self, service: DocxService, docx_dir: Path
    ) -> None:
        text = _extract(service, docx_dir / "tracked_en.docx").pages[0].text
        assert "The pump runs quietly at full load." in text
        assert "loudly" not in text

    def test_field_codes_are_not_text(self, service: DocxService, docx_dir: Path) -> None:
        text = _extract(service, docx_dir / "tracked_en.docx").pages[0].text
        assert "PAGE" not in text

    def test_content_controls_are_read(self, service: DocxService, docx_dir: Path) -> None:
        text = _extract(service, docx_dir / "tracked_en.docx").pages[0].text
        assert "Warranty terms are inside a content control." in text


class TestPages:
    def test_a_hard_page_break_starts_a_new_page(
        self, service: DocxService, docx_dir: Path
    ) -> None:
        pages = _extract(service, docx_dir / "manual_ko.docx").pages
        assert [page.page_number for page in pages] == [1, 2]

    def test_rendered_page_breaks_number_the_pages(
        self, service: DocxService, docx_dir: Path
    ) -> None:
        pages = _extract(service, docx_dir / "tracked_en.docx").pages
        assert [page.page_number for page in pages] == [1, 2, 3]
        assert pages[1].blocks[0].text == "Page two starts with this paragraph."

    def test_a_paragraph_that_straddles_a_break_stays_whole_on_its_first_page(
        self, service: DocxService, docx_dir: Path
    ) -> None:
        pages = _extract(service, docx_dir / "tracked_en.docx").pages
        assert "This paragraph straddles the second page break." in pages[1].text
        assert pages[2].text == "Page three holds the last paragraph."

    def test_a_document_with_no_page_markers_is_one_page(
        self, service: DocxService, tmp_path: Path
    ) -> None:
        import docx

        document = docx.Document()
        for _ in range(5):
            document.add_paragraph("A document written by a tool that records no pages.")
        path = tmp_path / "plain.docx"
        document.save(path)
        assert len(_extract(service, path).pages) == 1


class TestErrors:
    def test_an_encrypted_file_is_unsupported(self, service: DocxService, docx_dir: Path) -> None:
        with pytest.raises(ExtractionUnsupportedError, match="password-protected"):
            _extract(service, docx_dir / "locked.docx")

    def test_a_file_that_is_not_a_zip_is_an_error(
        self, service: DocxService, tmp_path: Path
    ) -> None:
        path = tmp_path / "fake.docx"
        path.write_bytes(b"just some text")
        with pytest.raises(ExtractionError):
            _extract(service, path)

    def test_a_zip_without_a_document_is_an_error(
        self, service: DocxService, tmp_path: Path
    ) -> None:
        import zipfile

        path = tmp_path / "other.docx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("hello.txt", "hi")
        with pytest.raises(ExtractionError):
            _extract(service, path)

    def test_an_empty_document_is_unsupported(self, service: DocxService, tmp_path: Path) -> None:
        import docx

        path = tmp_path / "empty.docx"
        docx.Document().save(path)
        with pytest.raises(ExtractionUnsupportedError):
            _extract(service, path)
