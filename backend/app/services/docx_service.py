# backend/app/services/docx_service.py
"""Word (.docx) extraction into the same pages and blocks a PDF produces.

The body is walked as raw WordprocessingML, in document order, because python-docx's
paragraph.text drops the things this needs: which text was deleted under tracked
changes, where Word last broke the page, and soft hyphens.

A .docx has no fixed pages. Pages are approximated from the markers Word leaves in
the file: hard page breaks, section breaks that start a new page, and the
w:lastRenderedPageBreak that Word writes on every save where it last laid out a page.
A file written by another tool may carry none of them and is then a single page.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from app.logging_config import get_logger
from app.models.domain import (
    BlockKind,
    DocumentMeta,
    ExtractedDocument,
    ExtractedPage,
    PageBlock,
)
from app.services.extractors import (
    ExtractionError,
    ExtractionUnsupportedError,
    table_markdown,
)
from app.services.text_service import detect_language, join_wrapped_lines, normalize_text

logger = get_logger("docx")

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ZIP_MAGIC = b"PK\x03\x04"
# An encrypted .docx is an OLE compound file, and so is a legacy .doc renamed .docx.
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _tag(name: str, namespace: str = _W) -> str:
    return f"{{{namespace}}}{name}"


_P, _TBL, _TR, _TC = _tag("p"), _tag("tbl"), _tag("tr"), _tag("tc")
_SDT, _SDT_CONTENT, _CUSTOM_XML = _tag("sdt"), _tag("sdtContent"), _tag("customXml")
_T, _TAB, _BR, _CR = _tag("t"), _tag("tab"), _tag("br"), _tag("cr")
_RENDERED_BREAK = _tag("lastRenderedPageBreak")
_SOFT_HYPHEN, _NO_BREAK_HYPHEN = _tag("softHyphen"), _tag("noBreakHyphen")
_PPR, _PSTYLE, _OUTLINE, _SECTPR, _TYPE = (
    _tag("pPr"),
    _tag("pStyle"),
    _tag("outlineLvl"),
    _tag("sectPr"),
    _tag("type"),
)
_VAL, _TYPE_ATTR = _tag("val"), _tag("type")
# Deleted text under tracked changes, moved-away text, the duplicate that
# AlternateContent carries for old readers, and properties that hold no text.
_SKIPPED = {
    _tag("del"),
    _tag("moveFrom"),
    _tag("delText"),
    _tag("instrText"),
    _tag("Fallback", _MC),
    _PPR,
    _tag("rPr"),
    _tag("footnoteReference"),
    _tag("endnoteReference"),
}
_NEW_PAGE_SECTIONS = {"nextPage", "oddPage", "evenPage"}
_HEADING_NAME = re.compile(r"^(heading [1-9]|title)$", re.IGNORECASE)
_HEADING_ID = re.compile(r"^(Heading[1-9]|Title)$")
_MAX_OUTLINE_LEVEL = 8
_LANGUAGE_SAMPLE_CHARS = 5000

_PAGE_BREAK = object()  # a marker in a paragraph's stream of text pieces


class DocxService:
    suffixes = (".docx",)
    media_type = DOCX_MEDIA_TYPE

    def __init__(self, min_document_chars: int = 20, extract_tables: bool = True) -> None:
        self._min_document_chars = min_document_chars
        self._extract_tables = extract_tables

    def extract(self, path: Path, *, document_id: str, file_hash: str) -> ExtractedDocument:
        self._check_container(path)
        try:
            import docx

            document = docx.Document(str(path))
        except Exception as exc:  # PackageNotFoundError, BadZipFile, KeyError, XMLSyntaxError
            raise ExtractionError(f"cannot open {path.name}: {exc}") from exc

        headings = self._heading_styles(document.styles.element)
        builder = _PageBuilder()
        try:
            for element in self._body_items(document.element.body):
                if element.tag == _P:
                    self._add_paragraph(element, headings, builder)
                else:
                    self._add_table(element, builder)
        except Exception as exc:
            raise ExtractionError(f"failed to read {path.name}: {exc}") from exc

        pages = tuple(
            ExtractedPage(
                page_number=index + 1,
                blocks=tuple(blocks),
                text=normalize_text("\n\n".join(block.text for block in blocks)),
            )
            for index, blocks in enumerate(builder.pages())
        )
        title = (document.core_properties.title or "").strip() or None
        meta = DocumentMeta(
            document_id=document_id,
            filename=path.name,
            filepath=str(path),
            folder=str(path.parent),
            file_hash=file_hash,
            modified_at=datetime.fromtimestamp(path.stat().st_mtime, tz=UTC),
            language=None,
            title=title,
        )
        extracted = ExtractedDocument(meta=meta, pages=pages)
        if extracted.char_count < self._min_document_chars:
            raise ExtractionUnsupportedError(f"{path.name} has no text to index")

        sample = "\n".join(page.text for page in pages)[:_LANGUAGE_SAMPLE_CHARS]
        language = detect_language(sample)
        logger.info(
            "extracted %s: pages=%d chars=%d language=%s",
            path.name,
            len(pages),
            extracted.char_count,
            language,
        )
        return ExtractedDocument(meta=replace(meta, language=language), pages=pages)

    # ------------------------------------------------------------- container

    @staticmethod
    def _check_container(path: Path) -> None:
        try:
            with path.open("rb") as handle:
                head = handle.read(8)
        except OSError as exc:
            raise ExtractionError(f"cannot open {path.name}: {exc}") from exc
        if head.startswith(OLE_MAGIC):
            raise ExtractionUnsupportedError(
                f"{path.name} is password-protected, or an old .doc file renamed to .docx"
            )
        if not head.startswith(ZIP_MAGIC):
            raise ExtractionError(f"{path.name} is not a Word .docx file")
        try:
            with zipfile.ZipFile(path) as archive:
                archive.getinfo("word/document.xml")
        except (zipfile.BadZipFile, KeyError) as exc:
            raise ExtractionError(f"{path.name} is not a Word .docx file") from exc

    # ---------------------------------------------------------------- styles

    @staticmethod
    def _heading_styles(styles: Any) -> set[str]:
        """Paragraph style ids that are headings, following basedOn chains.

        The outline level is what Word itself uses, and unlike the display name it is
        the same in every language: Korean Word calls Heading 1 "제목 1".
        """
        based_on: dict[str, str] = {}
        direct: set[str] = set()
        for style in styles.iterchildren(_tag("style")):
            if style.get(_TYPE_ATTR) != "paragraph":
                continue
            style_id = style.get(_tag("styleId")) or ""
            name = style.find(_tag("name"))
            name_value = name.get(_VAL) if name is not None else ""
            outline = style.find(f"{_PPR}/{_OUTLINE}")
            if (
                _HEADING_ID.match(style_id)
                or _HEADING_NAME.match(name_value or "")
                or (outline is not None and _outline_level(outline) <= _MAX_OUTLINE_LEVEL)
            ):
                direct.add(style_id)
            parent = style.find(_tag("basedOn"))
            if parent is not None and parent.get(_VAL):
                based_on[style_id] = parent.get(_VAL)

        headings = set(direct)
        for style_id in based_on:
            seen: set[str] = set()
            current: str | None = style_id
            while current is not None and current not in seen:
                if current in direct:
                    headings.add(style_id)
                    break
                seen.add(current)
                current = based_on.get(current)
        return headings

    # ------------------------------------------------------------------ body

    def _body_items(self, container: Any) -> list[Any]:
        """Paragraphs and tables in order, unwrapping content controls."""
        items: list[Any] = []
        for child in container.iterchildren():
            if child.tag in (_P, _TBL):
                items.append(child)
            elif child.tag == _SDT:
                content = child.find(_SDT_CONTENT)
                if content is not None:
                    items.extend(self._body_items(content))
            elif child.tag == _CUSTOM_XML:
                items.extend(self._body_items(child))
        return items

    def _add_paragraph(self, paragraph: Any, headings: set[str], builder: _PageBuilder) -> None:
        pieces: list[object] = []
        _collect(paragraph, pieces)
        text = _pieces_text(pieces)
        breaks = sum(1 for piece in pieces if piece is _PAGE_BREAK)
        leading = _leading_breaks(pieces)

        for _ in range(leading):
            builder.page_break()
        if text:
            builder.add(self._paragraph_kind(paragraph, headings), text)
        for _ in range(breaks - leading):
            builder.page_break()
        if _starts_new_page_after(paragraph):
            builder.page_break()

    @staticmethod
    def _paragraph_kind(paragraph: Any, headings: set[str]) -> BlockKind:
        properties = paragraph.find(_PPR)
        if properties is not None:
            outline = properties.find(_OUTLINE)
            if outline is not None:
                return "heading" if _outline_level(outline) <= _MAX_OUTLINE_LEVEL else "paragraph"
            style = properties.find(_PSTYLE)
            if style is not None and style.get(_VAL) in headings:
                return "heading"
        return "paragraph"

    def _add_table(self, table: Any, builder: _PageBuilder) -> None:
        rows: list[list[str]] = []
        breaks = 0
        for row in table.iterchildren(_TR):
            cells: list[str] = []
            for cell in row.iterchildren(_TC):
                texts: list[str] = []
                for paragraph in cell.iter(_P):  # nested tables flatten into the cell
                    pieces: list[object] = []
                    _collect(paragraph, pieces)
                    breaks += sum(1 for piece in pieces if piece is _PAGE_BREAK)
                    if text := _pieces_text(pieces):
                        texts.append(text)
                cells.append(" ".join(texts))
            if any(cells):
                rows.append(cells)

        if self._extract_tables and len(rows) >= 2:
            builder.add("table", table_markdown(rows))
        else:
            # A one-row table is almost always layout, not data.
            for cells in rows:
                for text in cells:
                    if text:
                        builder.add("paragraph", text)
        if breaks:
            builder.page_break()


class _PageBuilder:
    """Collects blocks and never produces an empty page.

    Word often writes both a hard break and a rendered break for the same page
    boundary; a break on a page that has nothing on it yet is the same boundary.
    """

    def __init__(self) -> None:
        self._pages: list[list[PageBlock]] = [[]]

    def add(self, kind: BlockKind, text: str) -> None:
        text = normalize_text(text)
        if text:
            page = self._pages[-1]
            page.append(PageBlock(kind=kind, text=text, order=len(page)))

    def page_break(self) -> None:
        if self._pages[-1]:
            self._pages.append([])

    def pages(self) -> list[list[PageBlock]]:
        return [page for page in self._pages if page] or [[]]


def _collect(element: Any, pieces: list[object]) -> None:
    """Text pieces and page-break markers under one paragraph, in order."""
    for child in element.iterchildren():
        tag = child.tag
        if not isinstance(tag, str) or tag in _SKIPPED:
            continue  # comments and processing instructions have non-string tags
        if tag == _T:
            pieces.append(child.text or "")
        elif tag == _TAB:
            pieces.append(" ")
        elif tag == _BR:
            kind = child.get(_TYPE_ATTR)
            if kind == "page":
                pieces.append(_PAGE_BREAK)
            elif kind != "column":
                pieces.append("\n")
        elif tag == _CR:
            pieces.append("\n")
        elif tag == _RENDERED_BREAK:
            pieces.append(_PAGE_BREAK)
        elif tag == _NO_BREAK_HYPHEN:
            pieces.append("-")
        elif tag == _SOFT_HYPHEN:
            continue
        else:
            _collect(child, pieces)


def _pieces_text(pieces: list[object]) -> str:
    raw = "".join(piece for piece in pieces if isinstance(piece, str))
    return join_wrapped_lines(raw.split("\n"))


def _leading_breaks(pieces: list[object]) -> int:
    """Breaks before any text: the paragraph itself starts on the next page.

    A break after some text leaves the paragraph on the page it started on, so a
    paragraph is never cut in two, and nor is a word Word happened to break at.
    """
    count = 0
    for piece in pieces:
        if piece is _PAGE_BREAK:
            count += 1
        elif isinstance(piece, str) and piece.strip():
            break
    return count


def _starts_new_page_after(paragraph: Any) -> bool:
    section = paragraph.find(f"{_PPR}/{_SECTPR}")
    if section is None:
        return False
    kind = section.find(_TYPE)
    value = kind.get(_VAL) if kind is not None else "nextPage"  # nextPage is the default
    return value in _NEW_PAGE_SECTIONS


def _outline_level(element: Any) -> int:
    try:
        return int(element.get(_VAL, "9"))
    except ValueError:
        return 9
